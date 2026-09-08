import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Background image or preview
    preview_img = r"C:\Users\HP\.gemini\antigravity-ide\brain\d6b0cd37-2352-47e0-8ad2-22d5f612c21e\slide_6_preview.png"
    if os.path.exists(preview_img):
        slide.shapes.add_picture(preview_img, 0, 0, width=prs.slide_width, height=prs.slide_height)

    pptx_path = r"C:\Users\HP\Downloads\TRACE_X_Live_Demo_and_Graph_Slide.pptx"
    prs.save(pptx_path)
    print("Saved PPTX to:", pptx_path)

if __name__ == "__main__":
    create_pptx()
